"""Kirim jadwal piket harian ke Telegram.

Dijalankan terjadwal di GitHub Actions (bebas dari laptop lokal):
  - Ambil jadwal dari Google Sheet (public, export CSV + XLSX via gid).
  - Susun pesan: piket semua shift (1, 2, 3) + PIC (sel hijau #00FF00).
  - Kirim via Telegram Bot API (sendMessage).

Env yang dibutuhkan:
  TELEGRAM_BOT_TOKEN  token bot Telegram
  TELEGRAM_CHAT_ID    chat ID tujuan (boleh angka, @username, atau "id1,id2")
  JADWAL_SHEET_URL    opsional, default = tab jadwal L1/L2&L3

Jalankan lokal:
  $env:TELEGRAM_BOT_TOKEN="..."; $env:TELEGRAM_CHAT_ID="903566652"
  python send_daily_jadwal.py --dry-run   (hanya tampilkan pesan)
  python send_daily_jadwal.py             (kirim sungguhan)
"""
import csv
import datetime
import io
import os
import re
import sys
import urllib.parse
import urllib.request
from zoneinfo import ZoneInfo

SHEET_URL = os.environ.get(
    "JADWAL_SHEET_URL",
    "https://docs.google.com/spreadsheets/d/"
    "1VlLQx0CxMeWStMjpkwQWiq8nbXaYqO2fJABAe7JNhmg/export?format=csv&gid=890618279",
)
TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID", "").strip()
# Untuk group forum mode: ID topik tujuan (message_thread_id)
THREAD_ID = os.environ.get("TELEGRAM_MESSAGE_THREAD_ID", "").strip()

SHIFTS = [
    ("Shift 1", "00:00", "09:00"),
    ("Shift 2", "09:00", "16:30"),
    ("Shift 3", "16:30", "23:59"),
]
SHIFT_START_M = {
    int(st.split(":")[0]) * 60 + int(st.split(":")[1]): (label, st, en)
    for label, st, en in SHIFTS
}

SHIFTS_L2 = [
    ("Shift 1", "08:00", "17:00"),
    ("Shift 2", "17:00", "24:00"),
]

SHIFTS_DEV = [
    ("Shift 1", "08:00", "17:00"),
    ("Shift 2", "17:00", "24:00"),
]

SHIFTS_L3 = [
    ("Shift 1", "08:00", "17:00"),
    ("Shift 2", "17:00", "24:00"),
]

NAMES_MAP = {
    "I NYOMAN WINDU": "WINDU",
    "RIZKY SENJA BUDIMAN": "RIZKY",
    "RIZKY SENJA": "RIZKY",
    "RIZKY": "RIZKY",
    # L2
    "ADAM (L2)": "ADAM",
    "ADAM": "ADAM",
    "RIYAN (L2)": "RIYAN",
    "RIYAN": "RIYAN",
    "ADYTAMA (L2)": "ADY",
    "ADYTAMA": "ADY",
    "ADY": "ADY",
    "FAISHAL (L2)": "FAISAL",
    "FAISHAL": "FAISAL",
    "FAISAL (L2)": "FAISAL",
    "FAISAL": "FAISAL",
    "ERWIN (L2)": "ERWIN",
    "ERWIN": "ERWIN",
    "RIZKY ARDIAN (L2)": "RIZKY ARDIAN",
    "RIZKY ARDIAN": "RIZKY ARDIAN",
    # DEV
    "REZA (DEV INSERA)": "REZA",
    "REZA": "REZA",
    "AJI (DEV BIMA)": "AJI",
    "AJI": "AJI",
    "WILDAN (DEV BIMA)": "WILDAN",
    "WILDAN": "WILDAN",
    # L3
    "KEVIN (L3)": "KEVIN",
    "KEVIN": "KEVIN",
    "GUNTUR (L3)": "GUNTUR",
    "GUNTUR": "GUNTUR",
    "DANA (L3)": "DANA",
    "DANA": "DANA",
}
USERNAMES = {
    # L1
    "ANNE": "@annemaswa27",
    "ALFIRA": "@hifirrr",
    "TYAS": "@ikhpani",
    "ROFINA": "@InaBiti",
    "RIZKA": "@rizkayln",
    "IHSAN": "@ihsanabdi123",
    "DIKA": "@dikaprndx",
    "WINDU": "@komangwindu",
    "RIZKY": "@rizkysenjab",
    # L2 & DEV
    "RIYAN": "@riyanalwandika",
    "ADAM": "@Adamilhamm",
    "ADY": "@ady_tama",
    "FAISAL": "@ishal223",
    "AJI": "@aifllvyy",
    "WILDAN": "@wildanfariz",
    "REZA": "@rezzamf",
}
GREEN_FILL = "FF00FF00"
SUMMARY_ROWS = {
    "SHIFT 1", "SHIFT 2", "SHIFT 3", "TOTAL MASUK PER HARI",
    "WANTIMPRES", "NAMA", "JAM KERJA",
}

DAY_NAMES = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
MONTH_NAMES = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
               "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

_MONTH_ABBR = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MEI": 5, "JUN": 6,
    "JUL": 7, "AGU": 8, "SEP": 9, "OKT": 10, "NOV": 11, "DES": 12,
}


def http_get(url, timeout=20):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def _month_num(word):
    """Nomor bulan dari nama Indonesia penuh/singkat (case-insensitive)."""
    key = re.sub(r'[^a-zA-Z]', '', word or '').upper()[:3]
    return _MONTH_ABBR.get(key)


def _parse_range_title(title, ref):
    """Parse judul periode tab, mis. '7 SEPTEMBER - 4 OKTOBER' atau
    '9 AGUSTUS - SEPTEMBER' (tanpa tanggal akhir = akhir bulan).

    Return (start_date, end_date) atau None jika tidak bisa diparse.
    """
    import calendar

    if not title:
        return None
    m = re.search(
        r"(\d{1,2})\s+([A-Za-z]+)\s*-\s*(?:(\d{1,2})\s+)?([A-Za-z]+)",
        str(title).strip(),
    )
    if not m:
        return None
    d1, w1, d2, w2 = int(m.group(1)), m.group(2), m.group(3), m.group(4)
    m1, m2 = _month_num(w1), _month_num(w2)
    if not m1 or not m2:
        return None
    try:
        start = datetime.date(ref.year, m1, d1)
    except ValueError:
        return None
    if d2:
        end = datetime.date(ref.year, m2, int(d2))
    else:
        end = datetime.date(ref.year, m2, calendar.monthrange(ref.year, m2)[1])
    if end < start:
        # Rentang lewat tahun baru, mis. Des-Jan.
        try:
            end = datetime.date(ref.year + 1, m2, end.day)
        except ValueError:
            return None
    if ref.date() < start:
        # Hari ini sebelum periode tahun ini -> periode tahun lalu.
        try:
            start = datetime.date(ref.year - 1, m1, d1)
            end = datetime.date(end.year - 1, end.month, end.day)
        except ValueError:
            return None
    return (start, end)


def _tab_title(rows):
    """Ambil judul periode dari baris-baris awal CSV (mis. '7 SEPTEMBER - 4 OKTOBER')."""
    for r in (rows or [])[:4]:
        for cell in r or []:
            if cell and re.search(r"\d{1,2}\s+[A-Za-z]+\s*-", str(cell)):
                return str(cell).strip()
    return ""


def resolve_sheet_url(now):
    """Deteksi tab sheet yang memuat tanggal target secara otomatis.

    Beberapa tab bisa overlap (mis. 'Agu-Sep' dan 'Sep-Okt' sama-sama
    memuat September). Jika begitu, pilih tab yang periodenya mencakup
    hari ini dengan tanggal mulai paling baru (jadwal yang berlaku).
    """
    if "JADWAL_SHEET_URL" in os.environ:
        return SHEET_URL
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", SHEET_URL)
    if not m:
        return SHEET_URL
    sheet_id = m.group(1)
    try:
        html = http_get(f"https://docs.google.com/spreadsheets/d/{sheet_id}/htmlview").decode("utf-8", errors="replace")
        tabs = re.findall(r'name:\s*"([^"]+)"[^\}]+?gid:\s*"(\d+)"', html)
        target_m = now.month
        target_d = now.day
        seen = set()
        candidates = []  # (csv_url, start_or_None)
        for name, gid in tabs:
            if gid in seen:
                continue
            seen.add(gid)
            clean = re.sub(r'[^a-zA-Z]', ' ', name).upper()
            matched_m = [v for k, v in _MONTH_ABBR.items() if any(t.startswith(k) for t in clean.split())]
            if target_m in matched_m:
                csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"
                csv_raw = http_get(csv_url).decode("utf-8", errors="replace")
                rows = list(csv.reader(io.StringIO(csv_raw)))
                for r in rows:
                    if r and r[0] == "NAMA":
                        if any(_cell_day(h, target_d) for h in r):
                            rng = _parse_range_title(_tab_title(rows), now)
                            if rng and rng[0] <= now.date() <= rng[1]:
                                candidates.append((csv_url, rng[0]))
                            else:
                                # Rentang tak terparse / tak mencakup: fallback
                                # (perilaku lama: tab cocok pertama menang).
                                candidates.append((csv_url, None))
                        break
        covering = [(u, s) for u, s in candidates if s is not None]
        if covering:
            covering.sort(key=lambda t: t[1], reverse=True)
            return covering[0][0]
        if candidates:
            return candidates[0][0]
    except Exception:
        pass
    return SHEET_URL


def _cell_day(value, day):
    if value is None:
        return False
    return str(value).strip().split(".")[0] == str(day)


def _slot_start_min(slot):
    try:
        s, _e = slot.split("-")
        sh, sm = (int(x) for x in s.strip().replace(".", ":").split(":"))
        return sh * 60 + sm
    except Exception:
        return None


def match_shift_slot(sheet_map, start_m):
    """Pilih key slot sheet untuk jam mulai shift config.

    Exact match dulu; jika sheet memakai label jam berbeda (mis. slot
    '08:00-17:00' untuk shift 09:00-16:30), ambil slot yang jam mulainya
    <= jam mulai shift (slot terakhir yang sudah mulai).
    """
    if not sheet_map:
        return None
    if start_m in sheet_map:
        return start_m
    starts = sorted(k for k in sheet_map if k <= start_m)
    return starts[-1] if starts else None


def parse_piket_csv(raw, day):
    """Parse CSV sheet -> {'L1': {slot_m: [tags]}, 'L2': {...}, 'DEV': {...}, 'L3': {...}}."""
    rows = list(csv.reader(io.StringIO(raw)))
    day_col = None
    for r in rows:
        if r and r[0] == "NAMA":
            for j, h in enumerate(r):
                if _cell_day(h, day):
                    day_col = j
                    break
            if day_col is not None:
                break
    if day_col is None:
        return None

    result = {"L1": {}, "L2": {}, "DEV": {}, "L3": {}}
    cur_section = "L1"
    cur_name = None

    for r in rows:
        if not r:
            continue
        c0 = (r[0] or "").strip()
        if c0:
            c0_u = c0.upper()
            if c0_u in ["L2", "JADWAL L2"]:
                cur_section = "L2"
                cur_name = None
                continue
            elif "DEV" in c0_u and "JADWAL" in c0_u:
                cur_section = "DEV"
                cur_name = None
                continue
            elif "L3" in c0_u and "JADWAL" in c0_u:
                cur_section = "L3"
                cur_name = None
                continue
            cur_name = c0

        if not cur_name or cur_name.upper() in SUMMARY_ROWS:
            continue

        slot = (r[1] or "").strip() if len(r) > 1 else ""
        slot_m = _slot_start_min(slot) if slot else None
        if slot_m is None or len(r) <= day_col:
            continue

        if (r[day_col] or "").strip() == "1":
            role = cur_section
            c_u = cur_name.upper()
            if "(L2)" in c_u:
                role = "L2"
            elif "DEV" in c_u:
                role = "DEV"
            elif "(L3)" in c_u:
                role = "L3"

            key = NAMES_MAP.get(c_u, c_u)
            tag = USERNAMES.get(key, cur_name)
            if tag and tag not in result[role].setdefault(slot_m, []):
                result[role][slot_m].append(tag)

    return result


def parse_pic_xlsx(raw, day):
    """PIC per role & shift (sel hijau) -> {'L1': {slot_m: tag}, 'L2': {...}, ...}."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active
    day_col = None
    for r in range(1, 40):
        if ws.cell(r, 1).value == "NAMA":
            for c in range(1, ws.max_column + 1):
                if _cell_day(ws.cell(r, c).value, day):
                    day_col = c
                    break
            if day_col is not None:
                break
    if day_col is None:
        return None

    result = {"L1": {}, "L2": {}, "DEV": {}, "L3": {}}
    cur_section = "L1"
    cur_name = None

    for r in range(1, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        c0 = str(c1).strip() if c1 else ""
        if c0:
            c0_u = c0.upper()
            if c0_u in ["L2", "JADWAL L2"]:
                cur_section = "L2"
                cur_name = None
                continue
            elif "DEV" in c0_u and "JADWAL" in c0_u:
                cur_section = "DEV"
                cur_name = None
                continue
            elif "L3" in c0_u and "JADWAL" in c0_u:
                cur_section = "L3"
                cur_name = None
                continue
            cur_name = c0

        if not cur_name or cur_name.upper() in SUMMARY_ROWS:
            continue

        c2 = ws.cell(r, 2).value
        slot = str(c2).strip() if c2 else ""
        slot_m = _slot_start_min(slot) if slot else None
        if slot_m is None:
            continue

        f = ws.cell(r, day_col).fill
        green = bool(f and f.patternType and f.fgColor and f.fgColor.rgb == GREEN_FILL)
        if green:
            role = cur_section
            c_u = cur_name.upper()
            if "(L2)" in c_u:
                role = "L2"
            elif "DEV" in c_u:
                role = "DEV"
            elif "(L3)" in c_u:
                role = "L3"

            key = NAMES_MAP.get(c_u, c_u)
            tag = USERNAMES.get(key, cur_name)
            if tag and slot_m not in result[role]:
                result[role][slot_m] = tag

    return result


def build_message(now, role="default"):
    day = now.day
    csv_url = resolve_sheet_url(now)
    xlsx_url = csv_url.replace("export?format=csv", "export?format=xlsx")
    piket = parse_piket_csv(http_get(csv_url).decode("utf-8", errors="replace"), day)
    pic = parse_pic_xlsx(
        http_get(xlsx_url), day
    )
    if piket is None:
        raise RuntimeError("Tanggal tidak ada di jadwal sheet / CSV tidak terbaca")

    title = "📋 Jadwal Piket Semua Role" if role == "all" else "📋 Jadwal Piket Hari Ini"
    lines = [
        title,
        f"{DAY_NAMES[now.weekday()]}, {now.day} {MONTH_NAMES[now.month - 1]} {now.year}",
        "",
        "🔹 Jadwal L1:",
    ]
    piket_l1 = piket.get("L1", {})
    pic_l1 = pic.get("L1", {}) if pic else {}
    for start_m in sorted(SHIFT_START_M):
        label, start, end = SHIFT_START_M[start_m]
        lines.append(f"👥 {label} ({start} - {end}):")
        key = match_shift_slot(piket_l1, start_m)
        tags = piket_l1.get(key) if key is not None else None
        if not tags:
            lines.append("• (tidak ada)")
        else:
            for t in tags:
                pic_mark = " — ⭐ PIC" if pic_l1 and key is not None and pic_l1.get(key) == t else ""
                lines.append(f"• {t}{pic_mark}")
        lines.append("")

    # L2
    piket_l2 = piket.get("L2", {})
    pic_l2 = pic.get("L2", {}) if pic else {}
    lines.append("🔹 Jadwal L2:")
    for label, start, end in SHIFTS_L2:
        lines.append(f"👥 {label} ({start} - {end}):")
        sh, sm = (int(x) for x in start.split(":"))
        start_m = sh * 60 + sm
        key = match_shift_slot(piket_l2, start_m)
        tags = piket_l2.get(key) if key is not None else None
        if not tags:
            lines.append("• (tidak ada)")
        else:
            for t in tags:
                pic_mark = " — ⭐ PIC" if pic_l2 and key is not None and pic_l2.get(key) == t else ""
                lines.append(f"• {t}{pic_mark}")
        lines.append("")

    if role == "all":
        # DEV
        piket_dev = piket.get("DEV", {})
        pic_dev = pic.get("DEV", {}) if pic else {}
        lines.append("🔹 Jadwal DEV (Insera & Bima):")
        has_dev = False
        for label, start, end in SHIFTS_DEV:
            sh, sm = (int(x) for x in start.split(":"))
            start_m = sh * 60 + sm
            key = match_shift_slot(piket_dev, start_m)
            tags = piket_dev.get(key) if key is not None else None
            if tags:
                has_dev = True
                lines.append(f"👥 {label} ({start} - {end}):")
                for t in tags:
                    pic_mark = " — ⭐ PIC" if pic_dev and key is not None and pic_dev.get(key) == t else ""
                    lines.append(f"• {t}{pic_mark}")
        if not has_dev:
            lines.append("• (tidak ada)")
        lines.append("")

        # L3
        piket_l3 = piket.get("L3", {})
        pic_l3 = pic.get("L3", {}) if pic else {}
        lines.append("🔹 Jadwal L3:")
        has_l3 = False
        for label, start, end in SHIFTS_L3:
            sh, sm = (int(x) for x in start.split(":"))
            start_m = sh * 60 + sm
            key = match_shift_slot(piket_l3, start_m)
            tags = piket_l3.get(key) if key is not None else None
            if tags:
                has_l3 = True
                lines.append(f"👥 {label} ({start} - {end}):")
                for t in tags:
                    pic_mark = " — ⭐ PIC" if pic_l3 and key is not None and pic_l3.get(key) == t else ""
                    lines.append(f"• {t}{pic_mark}")
        if not has_l3:
            lines.append("• (tidak ada)")
        lines.append("")

    lines.append("Sumber: Google Sheet jadwal")
    return "\n".join(lines)


def send_message(text):
    url = f"https://api.telegram.org/bot{TOKEN}/sendMessage"
    for cid in CHAT_ID.split(","):
        cid = cid.strip()
        if not cid:
            continue
        params = {"chat_id": cid, "text": text, "disable_web_page_preview": "true"}
        if THREAD_ID:
            params["message_thread_id"] = THREAD_ID
        data = urllib.parse.urlencode(params).encode()
        req = urllib.request.Request(url, data=data)
        with urllib.request.urlopen(req, timeout=30) as resp:
            body = resp.read().decode()
            if '"ok":false' in body or resp.status != 200:
                raise RuntimeError(f"Telegram gagal ke {cid}: {body[:200]}")


def main():
    role = "all" if ("--all" in sys.argv or "--role" in sys.argv and "all" in sys.argv) else "default"
    if "--dry-run" in sys.argv:
        now = datetime.datetime.now(ZoneInfo("Asia/Jakarta"))
        print(build_message(now, role=role))
        return 0
    if not TOKEN or not CHAT_ID:
        print("TELEGRAM_BOT_TOKEN / TELEGRAM_CHAT_ID belum diatur", file=sys.stderr)
        return 2
    now = datetime.datetime.now(ZoneInfo("Asia/Jakarta"))
    try:
        text = build_message(now, role=role)
    except Exception as e:
        text = f"⚠️ Gagal mengambil jadwal piket: {e}"
        try:
            send_message(text)
        except Exception:
            print(text, file=sys.stderr)
        return 1
    send_message(text)
    print("Terkirim ke:", CHAT_ID)
    return 0


if __name__ == "__main__":
    sys.exit(main())
